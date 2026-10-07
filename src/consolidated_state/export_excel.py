from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pandas as pd
from openpyxl.styles import Font

from .models import StatementSummary

DISCLAIMER = "Consolidated report derived from source statements. Not an official bank-issued statement."
MONEY_FORMAT = '$#,##0.00;[Red]-$#,##0.00'


def _sum_money(series: pd.Series) -> Decimal:
    total = Decimal("0.00")
    for value in series:
        if value is None or pd.isna(value):
            continue
        total += value if isinstance(value, Decimal) else Decimal(str(value))
    return total


def _overview(
    transactions: pd.DataFrame,
    statements: list[StatementSummary],
) -> pd.DataFrame:
    credits = transactions[transactions["credit"].notna()] if not transactions.empty else transactions
    debits = transactions[transactions["debit"].notna()] if not transactions.empty else transactions

    total_credits = _sum_money(credits["credit"]) if not credits.empty else Decimal("0.00")
    total_debits = _sum_money(debits["debit"]) if not debits.empty else Decimal("0.00")

    dated_statements = [
        statement
        for statement in statements
        if statement.statement_start and statement.statement_end
    ]
    if dated_statements:
        covered_period = (
            f"{min(statement.statement_start for statement in dated_statements).isoformat()} to "
            f"{max(statement.statement_end for statement in dated_statements).isoformat()}"
        )
    elif transactions.empty:
        covered_period = ""
    else:
        dates = pd.to_datetime(transactions["date"])
        covered_period = f"{dates.min().date().isoformat()} to {dates.max().date().isoformat()}"

    currencies = {
        statement.currency
        for statement in statements
        if statement.currency
    }
    if len(currencies) == 1:
        currency = next(iter(currencies))
    elif len(currencies) > 1:
        currency = "MIXED"
    else:
        currency = "Not explicitly identified"

    ordered_statements = sorted(
        statements,
        key=lambda statement: statement.statement_period or "9999",
    )
    opening = next(
        (
            statement.opening_balance
            for statement in ordered_statements
            if statement.opening_balance is not None
        ),
        None,
    )
    closing = next(
        (
            statement.closing_balance
            for statement in reversed(ordered_statements)
            if statement.closing_balance is not None
        ),
        None,
    )

    rows = [
        ("Covered period", covered_period),
        ("Currency", currency),
        ("Source statements", len(statements)),
        ("Total transactions", len(transactions)),
        ("Deposits / additions count", len(credits)),
        ("Deposits / additions total", total_credits),
        ("Withdrawals / subtractions count", len(debits)),
        ("Withdrawals / subtractions total", total_debits),
        ("Net change", total_credits - total_debits),
        ("Beginning balance", opening),
        ("Ending balance", closing),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def build_excel_report(
    transactions: pd.DataFrame,
    summary: pd.DataFrame,
    validation: pd.DataFrame,
    statements: list[StatementSummary] | None = None,
) -> bytes:
    statements = statements or []
    buffer = BytesIO()

    credits = transactions[transactions["credit"].notna()].copy()
    debits = transactions[transactions["debit"].notna()].copy()

    credit_columns = ["date", "description", "credit", "source_file", "statement_period"]
    debit_columns = ["date", "description", "debit", "source_file", "statement_period"]

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _overview(transactions, statements).to_excel(writer, sheet_name="Overview", index=False)
        credits[credit_columns].to_excel(writer, sheet_name="Credits & Deposits", index=False)
        debits[debit_columns].to_excel(writer, sheet_name="Debits & Withdrawals", index=False)
        transactions.to_excel(writer, sheet_name="All Transactions", index=False)
        summary.to_excel(writer, sheet_name="Monthly Summary", index=False)
        validation.to_excel(writer, sheet_name="Validation", index=False)

        workbook = writer.book
        info = workbook.create_sheet("About")
        info["A1"] = "ConsolidatedState"
        info["A1"].font = Font(bold=True, size=14)
        info["A3"] = DISCLAIMER
        info.column_dimensions["A"].width = 95

        data_sheets = [
            "Overview",
            "Credits & Deposits",
            "Debits & Withdrawals",
            "All Transactions",
            "Monthly Summary",
            "Validation",
        ]

        for sheet_name in data_sheets:
            sheet = workbook[sheet_name]
            sheet.freeze_panes = "A2"

            for cell in sheet[1]:
                cell.font = Font(bold=True)

            for column_cells in sheet.columns:
                max_len = max(
                    (len(str(cell.value)) if cell.value is not None else 0)
                    for cell in column_cells
                )
                sheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(max_len + 2, 10),
                    48,
                )

        for sheet_name, money_headers in {
            "Credits & Deposits": ("credit",),
            "Debits & Withdrawals": ("debit",),
            "All Transactions": ("debit", "credit", "balance"),
            "Monthly Summary": ("total_debits", "total_credits", "net"),
        }.items():
            sheet = workbook[sheet_name]
            headers = {cell.value: cell.column for cell in sheet[1]}

            for header in money_headers:
                if header not in headers:
                    continue
                for column_cells in sheet.iter_cols(
                    min_col=headers[header],
                    max_col=headers[header],
                    min_row=2,
                ):
                    for cell in column_cells:
                        cell.number_format = MONEY_FORMAT

        overview_sheet = workbook["Overview"]
        money_metrics = {
            "Deposits / additions total",
            "Withdrawals / subtractions total",
            "Net change",
            "Beginning balance",
            "Ending balance",
        }
        for row in range(2, overview_sheet.max_row + 1):
            if overview_sheet.cell(row=row, column=1).value in money_metrics:
                overview_sheet.cell(row=row, column=2).number_format = MONEY_FORMAT

    return buffer.getvalue()
