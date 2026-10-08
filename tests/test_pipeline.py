from datetime import date
from decimal import Decimal
from io import BytesIO

import fitz
import pandas as pd
from openpyxl import load_workbook

from consolidated_state.consolidate import (
    currency_summary,
    find_duplicate_transactions,
    monthly_summary,
    transactions_to_frame,
)
from consolidated_state.export_excel import build_excel_report
from consolidated_state.export_pdf import build_pdf_report
from consolidated_state.models import StatementSummary, Transaction
from consolidated_state.validate import (
    reconcile_statement,
    validate_statement_set,
    validate_transaction_rows,
)


def sample_transactions():
    return [
        Transaction(
            date(2026, 1, 2),
            "Deposit",
            credit=Decimal("100.00"),
            source_file="jan.pdf",
            statement_period="2026-01",
            currency="USD",
        ),
        Transaction(
            date(2026, 1, 3),
            "Coffee",
            debit=Decimal("5.00"),
            source_file="jan.pdf",
            statement_period="2026-01",
            currency="USD",
        ),
    ]


def test_monthly_summary_and_reconciliation():
    frame = transactions_to_frame(sample_transactions())
    summary = monthly_summary(frame)
    assert summary.iloc[0]["currency"] == "USD"
    assert summary.iloc[0]["total_debits"] == Decimal("5.00")
    assert summary.iloc[0]["total_credits"] == Decimal("100.00")
    assert summary.iloc[0]["net"] == Decimal("95.00")

    statement = StatementSummary("jan.pdf", "2026-01", Decimal("1000.00"), Decimal("1095.00"))
    result = reconcile_statement(statement, frame)
    assert result["status"] == "PASS"
    assert result["difference"] == Decimal("0.00")


def test_currency_summary_keeps_unknown_sources_separate():
    transactions = [
        Transaction(date(2026, 1, 1), "A", credit=Decimal("10.00"), source_file="a.pdf"),
        Transaction(date(2026, 1, 2), "B", credit=Decimal("20.00"), source_file="b.pdf"),
    ]
    summary = currency_summary(transactions_to_frame(transactions))
    assert len(summary) == 2
    assert set(summary["currency"]) == {"Unknown (a.pdf)", "Unknown (b.pdf)"}


def test_decimal_math_is_exact_at_cent_level():
    transactions = [
        Transaction(date(2026, 1, 1), "A", debit=Decimal("0.10"), source_file="jan.pdf", currency="USD"),
        Transaction(date(2026, 1, 2), "B", debit=Decimal("0.20"), source_file="jan.pdf", currency="USD"),
        Transaction(date(2026, 1, 3), "C", credit=Decimal("0.30"), source_file="jan.pdf", currency="USD"),
    ]
    summary = monthly_summary(transactions_to_frame(transactions))
    assert summary.iloc[0]["total_debits"] == Decimal("0.30")
    assert summary.iloc[0]["total_credits"] == Decimal("0.30")
    assert summary.iloc[0]["net"] == Decimal("0.00")


def test_transaction_integrity_accepts_valid_rows():
    result = validate_transaction_rows(transactions_to_frame(sample_transactions()))
    assert result["status"] == "PASS"


def test_transaction_integrity_rejects_both_debit_and_credit():
    bad = Transaction(
        date(2026, 1, 4),
        "Invalid row",
        debit=Decimal("1.00"),
        credit=Decimal("1.00"),
        source_file="jan.pdf",
    )
    result = validate_transaction_rows(transactions_to_frame([bad]))
    assert result["status"] == "FAIL"
    assert "invalid_amount_sides=1" in result["detail"]


def test_transaction_integrity_rejects_negative_amount():
    bad = Transaction(
        date(2026, 1, 4),
        "Invalid row",
        debit=Decimal("-1.00"),
        source_file="jan.pdf",
    )
    result = validate_transaction_rows(transactions_to_frame([bad]))
    assert result["status"] == "FAIL"
    assert "negative_amount=1" in result["detail"]


def test_duplicate_detection_across_same_detected_account():
    first = Transaction(date(2026, 1, 2), "Same", debit=Decimal("10.00"), source_file="a.pdf")
    second = Transaction(date(2026, 1, 2), "Same", debit=Decimal("10.00"), source_file="b.pdf")
    frame = transactions_to_frame([first, second])
    statements = [
        StatementSummary("a.pdf", account_fingerprint="same"),
        StatementSummary("b.pdf", account_fingerprint="same"),
    ]
    assert len(find_duplicate_transactions(frame, statements)) == 2


def test_identical_transactions_across_different_accounts_are_not_duplicates():
    first = Transaction(date(2026, 1, 2), "Same", debit=Decimal("10.00"), source_file="a.pdf")
    second = Transaction(date(2026, 1, 2), "Same", debit=Decimal("10.00"), source_file="b.pdf")
    frame = transactions_to_frame([first, second])
    statements = [
        StatementSummary("a.pdf", account_fingerprint="acct-a"),
        StatementSummary("b.pdf", account_fingerprint="acct-b"),
    ]
    assert find_duplicate_transactions(frame, statements).empty


def test_repeated_transaction_inside_one_statement_is_not_automatically_duplicate():
    first = Transaction(date(2026, 1, 2), "Same", debit=Decimal("10.00"), balance=Decimal("90.00"), source_file="a.pdf")
    second = Transaction(date(2026, 1, 2), "Same", debit=Decimal("10.00"), balance=Decimal("80.00"), source_file="a.pdf")
    frame = transactions_to_frame([first, second])
    assert find_duplicate_transactions(frame).empty


def test_statement_set_allows_multiple_accounts_and_currencies():
    statements = [
        StatementSummary(
            "a.pdf",
            "2026-01-01 to 2026-01-31",
            Decimal("100.00"),
            Decimal("110.00"),
            statement_start=date(2026, 1, 1),
            statement_end=date(2026, 1, 31),
            account_fingerprint="acct-a",
            currency="USD",
        ),
        StatementSummary(
            "b.pdf",
            "2026-02-01 to 2026-02-28",
            Decimal("500.00"),
            Decimal("520.00"),
            statement_start=date(2026, 2, 1),
            statement_end=date(2026, 2, 28),
            account_fingerprint="acct-b",
            currency="EUR",
        ),
    ]
    rows = validate_statement_set(statements)
    statuses = {row["check"]: row["status"] for row in rows}
    assert statuses["account_consistency"] == "WARNING"
    assert statuses["currency_consistency"] == "WARNING"
    assert statuses["statement_period_overlap"] == "SKIPPED"
    assert statuses["statement_period_gaps"] == "SKIPPED"
    assert statuses["balance_continuity"] == "SKIPPED"


def test_statement_set_detects_overlap_within_same_account():
    statements = [
        StatementSummary(
            "a.pdf",
            "2026-01-01 to 2026-01-31",
            Decimal("100.00"),
            Decimal("110.00"),
            statement_start=date(2026, 1, 1),
            statement_end=date(2026, 1, 31),
            account_fingerprint="same",
            currency="USD",
        ),
        StatementSummary(
            "b.pdf",
            "2026-01-31 to 2026-02-28",
            Decimal("999.00"),
            Decimal("120.00"),
            statement_start=date(2026, 1, 31),
            statement_end=date(2026, 2, 28),
            account_fingerprint="same",
            currency="USD",
        ),
    ]
    rows = validate_statement_set(statements)
    statuses = {row["check"]: row["status"] for row in rows}
    assert statuses["statement_period_overlap"] == "NEEDS_REVIEW"


def test_statement_set_detects_adjacent_balance_discontinuity():
    statements = [
        StatementSummary(
            "a.pdf",
            "2026-01-01 to 2026-01-31",
            Decimal("100.00"),
            Decimal("110.00"),
            statement_start=date(2026, 1, 1),
            statement_end=date(2026, 1, 31),
            account_fingerprint="same",
            currency="USD",
        ),
        StatementSummary(
            "b.pdf",
            "2026-02-01 to 2026-02-28",
            Decimal("111.00"),
            Decimal("120.00"),
            statement_start=date(2026, 2, 1),
            statement_end=date(2026, 2, 28),
            account_fingerprint="same",
            currency="USD",
        ),
    ]
    rows = validate_statement_set(statements)
    statuses = {row["check"]: row["status"] for row in rows}
    assert statuses["balance_continuity"] == "FAIL"


def test_excel_contains_expected_sheets():
    transactions = transactions_to_frame(sample_transactions())
    summary = monthly_summary(transactions)
    validation = pd.DataFrame([{"source_file": "jan.pdf", "check": "test", "status": "PASS", "detail": "ok", "difference": Decimal("0.00")}])
    statement = StatementSummary(
        "jan.pdf",
        "2026-01-01 to 2026-01-31",
        Decimal("1000.00"),
        Decimal("1095.00"),
        statement_start=date(2026, 1, 1),
        statement_end=date(2026, 1, 31),
        account_fingerprint="same",
        currency="USD",
        layout_strategy="ledger-two-sided-columns-v2",
    )
    data = build_excel_report(transactions, summary, validation, [statement])
    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == [
        "Overview",
        "Currency Summary",
        "Source Statements",
        "Credits & Deposits",
        "Debits & Withdrawals",
        "All Transactions",
        "Monthly Summary",
        "Validation",
        "About",
    ]

    source_statements = workbook["Source Statements"]
    assert source_statements.cell(row=2, column=8).value == "ledger-two-sided-columns-v2"

    currency_sheet = workbook["Currency Summary"]
    assert currency_sheet.cell(row=2, column=1).value == "USD"

    overview = workbook["Overview"]
    metrics = {
        overview.cell(row=row, column=1).value: overview.cell(row=row, column=2).value
        for row in range(2, overview.max_row + 1)
    }
    assert metrics["Currency handling"] == "USD"
    assert metrics["Covered period"] == "2026-01-01 to 2026-01-31"
    assert metrics["Deposits / additions count"] == 1
    assert metrics["Withdrawals / subtractions count"] == 1
    assert metrics["Ending balance"] == 1095


def test_pdf_report_is_valid_and_contains_expected_sections():
    transactions = transactions_to_frame(sample_transactions())
    summary = monthly_summary(transactions)
    validation = pd.DataFrame(
        [
            {
                "source_file": "jan.pdf",
                "check": "test",
                "status": "PASS",
                "detail": "ok",
                "difference": Decimal("0.00"),
            }
        ]
    )
    statement = StatementSummary(
        "jan.pdf",
        "2026-01-01 to 2026-01-31",
        Decimal("1000.00"),
        Decimal("1095.00"),
        statement_start=date(2026, 1, 1),
        statement_end=date(2026, 1, 31),
        account_fingerprint="same",
        currency="USD",
        layout_strategy="ledger-two-sided-columns-v2",
    )

    data = build_pdf_report(
        transactions,
        summary,
        validation,
        [statement],
    )

    assert data.startswith(b"%PDF")
    with fitz.open(stream=data, filetype="pdf") as document:
        assert document.page_count >= 1
        text = "\n".join(page.get_text("text") for page in document)

    assert "Consolidated Bank Statement Report" in text
    assert "Currency Summary" in text
    assert "Credits / Deposits / Additions" in text
    assert "Debits / Withdrawals / Subtractions" in text
    assert "Validation" in text
