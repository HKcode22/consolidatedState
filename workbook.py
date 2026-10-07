from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


@dataclass(frozen=True)
class Transaction:
    transaction_date: date
    description: str
    debit: Decimal | None
    credit: Decimal | None
    balance: Decimal | None
    source_filename: str
    statement_period: str


VALIDATION_COLUMNS = ("Check", "Status", "Details", "Source File", "Statement Period")


def _as_number(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _format_sheet(sheet, currency_columns: set[int] | None = None) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="244B62")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    sheet.freeze_panes = "A2"
    if sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        column_letter = column[0].column_letter
        max_length = max(
            (len(str(cell.value)) for cell in column if cell.value is not None),
            default=0,
        )
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)

    for column_index in currency_columns or set():
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=column_index).number_format = '#,##0.00;[Red](#,##0.00)'


def build_workbook(
    transactions: Sequence[Transaction],
    validation_rows: Sequence[Mapping[str, str]],
) -> bytes:
    """Build the requested three-sheet workbook entirely in memory."""
    workbook = Workbook()

    transaction_sheet = workbook.active
    transaction_sheet.title = "Transactions"
    transaction_sheet.append(
        [
            "Date",
            "Description",
            "Debit",
            "Credit",
            "Balance",
            "Source File",
            "Statement Period",
        ]
    )

    monthly_totals: dict[str, dict[str, Decimal | int]] = defaultdict(
        lambda: {
            "transaction_count": 0,
            "debits": Decimal("0"),
            "credits": Decimal("0"),
        }
    )

    for transaction in sorted(transactions, key=lambda item: item.transaction_date):
        transaction_sheet.append(
            [
                transaction.transaction_date,
                transaction.description,
                _as_number(transaction.debit),
                _as_number(transaction.credit),
                _as_number(transaction.balance),
                transaction.source_filename,
                transaction.statement_period,
            ]
        )
        month = transaction.transaction_date.strftime("%Y-%m")
        totals = monthly_totals[month]
        totals["transaction_count"] += 1
        if transaction.debit is not None:
            totals["debits"] += transaction.debit
        if transaction.credit is not None:
            totals["credits"] += transaction.credit

    for row in range(2, transaction_sheet.max_row + 1):
        transaction_sheet.cell(row=row, column=1).number_format = "yyyy-mm-dd"
    _format_sheet(transaction_sheet, {3, 4, 5})

    summary_sheet = workbook.create_sheet("Monthly Summary")
    summary_sheet.append(
        ["Month", "Transactions", "Total Debits", "Total Credits", "Net Activity"]
    )
    for month in sorted(monthly_totals):
        totals = monthly_totals[month]
        debits = totals["debits"]
        credits = totals["credits"]
        summary_sheet.append(
            [
                month,
                totals["transaction_count"],
                float(debits),
                float(credits),
                float(credits - debits),
            ]
        )
    _format_sheet(summary_sheet, {3, 4, 5})

    validation_sheet = workbook.create_sheet("Validation")
    validation_sheet.append(list(VALIDATION_COLUMNS))
    for item in validation_rows:
        validation_sheet.append([item.get(column, "") for column in VALIDATION_COLUMNS])
    _format_sheet(validation_sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
